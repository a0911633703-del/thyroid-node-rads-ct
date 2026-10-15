"""Apply saved full-development coefficients to the separate supplement workbook."""
from pathlib import Path
import argparse
import csv
import json
from openpyxl import load_workbook
from predict_from_parameters import predict_row


def main():
    here = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path, default=here.parent / 'Supplementary Data.xlsx')
    parser.add_argument('--input', type=Path, help='Alternative predictor CSV.')
    parser.add_argument('--parameters', type=Path, default=here / 'parameters/model_parameters.json')
    parser.add_argument('--output', type=Path, default=here / 'outputs/predicted_probabilities.csv')
    parser.add_argument('--models', help='Comma-separated saved names. Default: all five models.')
    parser.add_argument('--original-units', action='store_true')
    args = parser.parse_args()
    source = (args.input or args.workbook).resolve()
    parameters = args.parameters.resolve()
    output = args.output.resolve()
    if output in {source, parameters} or output.exists():
        parser.error('Choose a new output path that does not overwrite an existing file.')
    if not source.is_file():
        parser.error('Input not found. Set --workbook to Supplementary Data.xlsx or --input to a CSV.')
    if args.input:
        with source.open(encoding='utf-8-sig', newline='') as handle:
            rows = list(csv.DictReader(handle))
    else:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            values = workbook['Nodes'].iter_rows(values_only=True)
            fields = list(next(values))
            rows = [dict(zip(fields, values)) for values in values if any(x is not None for x in values)]
        finally:
            workbook.close()
        # The fitted coefficients describe the 248-node primary cohort.
        rows = [row for row in rows if int(row['primary_included']) == 1]
    saved = json.loads(parameters.read_text(encoding='utf-8'))
    names = args.models.split(',') if args.models else list(saved['models'])
    if set(names) - set(saved['models']):
        parser.error('Unknown model name.')
    records = []
    for row in rows:
        records.append({'node_id': row.get('node_id', ''), 'group_id': row.get('group_id', ''),
                        **{name + '_probability': format(predict_row(row, saved['models'][name], args.original_units), '.17g')
                           for name in names}})
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=['node_id', 'group_id'] + [n + '_probability' for n in names])
        writer.writeheader()
        writer.writerows(records)
    print(f'Saved full-development probabilities for {len(records)} nodes: {output}')


if __name__ == '__main__':
    main()
