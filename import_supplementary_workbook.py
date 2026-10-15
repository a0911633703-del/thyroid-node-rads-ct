"""Import the separate supplement into a NEW workspace for the analysis sources."""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import shutil
from openpyxl import load_workbook


def write_csv(path, rows, fields=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=fields or list(rows[0]), extrasaction='ignore')
        writer.writeheader()
        writer.writerows(rows)


def main():
    here = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path, default=here.parent / 'Supplementary Data.xlsx')
    parser.add_argument('--workspace', type=Path, default=here.parent / 'Analysis Workspace')
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    target = args.workspace.resolve()
    if not workbook_path.is_file():
        parser.error('Workbook not found. Set --workbook to Supplementary Data.xlsx.')
    if target.exists():
        parser.error('Workspace exists. Choose a new directory to preserve existing work.')
    if target == here or here in target.parents:
        parser.error('Workspace must be outside this code repository.')
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        def read_sheet(name):
            values = workbook[name].iter_rows(values_only=True)
            fields = list(next(values))
            return [dict(zip(fields, row)) for row in values if any(x is not None for x in row)]
        nodes = read_sheet('Nodes')
        predictions = read_sheet('OOF predictions')
    finally:
        workbook.close()
    for row in nodes:
        row['node_id'] = str(row['node_id'])
        row['group_id'] = str(row['group_id'])
    primary = [row for row in nodes if int(row['primary_included']) == 1]
    if len(nodes) != 254 or len(primary) != 248 or len({r['group_id'] for r in primary}) != 130:
        parser.error('Unexpected cohort counts.')
    if sum(int(r['outcome']) for r in primary) != 84:
        parser.error('Unexpected primary outcome count.')
    node_map = {r['node_id']: r for r in primary}
    models = {'NodeRADS', 'Imaging', 'Combined', 'RoutineCT', 'NodeRADS_RoutineCT'}
    if len(predictions) != 1240 or {r['model'] for r in predictions} != models:
        parser.error('Expected one OOF prediction per primary node and model.')
    seen = set()
    for row in predictions:
        node = node_map[str(row['node_id'])]
        key = (node['node_id'], row['model'])
        assert key not in seen
        seen.add(key)
        assert str(row['group_id']) == node['group_id']
        assert int(row['y_true']) == int(node['outcome'])
        assert int(row['fold']) == int(node['outer_fold'])
        assert int(row['prediction']) == int(float(row['probability']) >= float(row['threshold']))
        row['node_id'], row['group_id'], row['node_rads'] = node['node_id'], node['group_id'], node['node_rads']
    for row in nodes:
        row['patient_group_source'] = row['group_id']
        row['included_in_previous_249'] = bool(int(row['legacy_249_included']))
        for numerator, denominator, name in [('ce_ct', 'ce_ct_artery', 'ct_ratio_artery'),
                                             ('ce_ct', 'ce_ct_muscle', 'ct_ratio_muscle')]:
            row[name] = float(row[numerator]) / float(row[denominator]) if row[denominator] not in (None, '', 0) else None
        row['delta_ct'] = float(row['ce_ct']) - float(row['nc_ct'])
        row['delta_iodine'] = float(row['ce_iodine']) - float(row['nc_iodine'])
        row['spectral_slope_70_100'] = (float(row['ct_70kev']) - float(row['ct_100kev'])) / 30
        row['delta_casupp'] = float(row['casupp_70']) - float(row['casupp_30'])
    target.mkdir(parents=True)
    source_dir = target / 'source_measurements'
    # training.prepare() merges the two reference columns from this second file.
    source_fields = [k for k in nodes[0] if k not in {'ce_ct_artery', 'ce_ct_muscle'}]
    write_csv(source_dir / 'node_measurements_clean_v1.csv', nodes, source_fields)
    write_csv(source_dir / 'reference_and_legacy_variables_v1.csv', nodes,
              ['node_id', 'patient_group_source', 'ce_ct_artery', 'ce_ct_muscle'])
    analysis = target / 'analysis'
    write_csv(analysis / 'cohort_all.csv', nodes)
    write_csv(analysis / 'cohort_primary.csv', primary)
    write_csv(analysis / 'outer_fold_assignment.csv', nodes, ['node_id', 'group_id', 'outcome', 'outer_fold'])
    primary_names = {'NodeRADS', 'Imaging', 'Combined'}
    write_csv(analysis / 'primary_predictions.csv', [r for r in predictions if r['model'] in primary_names])
    revision = target / 'Revision_incremental_DECT_20261010'
    write_csv(revision / 'analysis/predictions_all_models.csv', predictions)
    write_csv(revision / 'analysis/new_model_predictions.csv', [r for r in predictions if r['model'] not in primary_names])
    for source in (here / 'analysis_source').glob('*.py'):
        folder = revision / 'code' if source.name == 'run_incremental_analysis.py' else target / 'code'
        folder.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, folder / source.name)
    shutil.copytree(here / 'saved_models/primary', analysis / 'models')
    shutil.copytree(here / 'saved_models/conventional_controls', revision / 'analysis/new_models')
    plan = json.loads((here / 'configuration/incremental_analysis_settings.json').read_text(encoding='utf-8'))
    plan['source_files_SHA256_before'] = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in source_dir.glob('*.csv')}
    (revision / 'analysis/incremental_analysis_plan.json').write_text(json.dumps(plan, indent=2) + '\n', encoding='utf-8')
    report = {'n_source_nodes': len(nodes), 'n_primary_nodes': len(primary), 'n_primary_patients': 130,
              'n_oof_predictions': len(predictions), 'analysis_run': False,
              'workbook_sha256': hashlib.sha256(workbook_path.read_bytes()).hexdigest()}
    (target / 'workbook_import_report.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
    print(f'Imported {len(nodes)} nodes and {len(predictions)} OOF predictions into {target}')


if __name__ == '__main__':
    main()
